"""Tests for the Excel export, including the audit sheet. All data is made up."""

from io import BytesIO

from openpyxl import load_workbook

from backend.app.main import AuditEntry, ExportRow, _build_workbook


def sheets(rows):
    return load_workbook(BytesIO(_build_workbook(rows)))


def row(**kw):
    return ExportRow(values={"first_name": "דנה", "id_number": "000000018"}, link="http://x/y.png", **kw)


def test_without_an_audit_there_is_only_the_data_sheet():
    assert sheets([row()]).sheetnames == ["Sheet"]


def test_the_audit_sheet_lists_every_field_with_its_journey():
    audit = [
        AuditEntry(field="שם פרטי", original="דנה", final="דנה", confidence=95, status="אוטומטי"),
        AuditEntry(field="תעודת זהות", original="00000018", final="000000018", confidence=60, status="תוקן ידנית"),
    ]
    wb = sheets([row(number=3, source="a.pdf · עמוד 2", audit=audit)])
    assert wb.sheetnames == ["Sheet", "יומן שינויים"]
    ws = wb["יומן שינויים"]
    assert [c.value for c in ws[1]] == ["שורה", "מקור", "שדה", "ערך שחולץ", "ערך סופי", "ציון ביטחון", "סטטוס"]
    assert [c.value for c in ws[2]] == [3, "a.pdf · עמוד 2", "שם פרטי", "דנה", "דנה", 95, "אוטומטי"]
    assert [c.value for c in ws[3]][3:] == ["00000018", "000000018", 60, "תוקן ידנית"]


def test_a_human_correction_is_shaded_and_the_rest_is_not():
    audit = [
        AuditEntry(field="א", original="1", final="1", status="אוטומטי"),
        AuditEntry(field="ב", original="1", final="2", status="תוקן ידנית"),
    ]
    ws = sheets([row(audit=audit)])["יומן שינויים"]
    assert ws["A3"].fill.fgColor.rgb.endswith("FDE9D2")
    assert ws["A2"].fill.fill_type is None


def test_a_missing_score_is_an_empty_cell():
    audit = [AuditEntry(field="א", original="", final="", confidence=None, status="לא נבדק (מסומן לבדיקה)")]
    assert sheets([row(audit=audit)])["יומן שינויים"]["F2"].value is None


def test_audit_values_that_look_like_formulas_stay_text():
    audit = [AuditEntry(field="א", original="=1+1", final="=2+2", status="תוקן ידנית")]
    ws = sheets([row(audit=audit)])["יומן שינויים"]
    assert ws["D2"].data_type == "s" and ws["D2"].value == "=1+1"
    assert ws["E2"].data_type == "s"
