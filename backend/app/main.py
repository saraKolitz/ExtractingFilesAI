import json
import re
from io import BytesIO
from pathlib import Path

from dotenv import load_dotenv
from fastapi import APIRouter, FastAPI, File, HTTPException, Request, Response, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel

# from .services.ai_service import ERROR_PREFIX, extract_data_from_image
# from .services.ai_service import extract_fields_with_gemini
# from .services.ai_service import extract_full_image_with_claude
from .services.ai_service import extract_with_agent_loop
# from .services.image_processing import extract_field_images
# from .services.ocr_service import get_word_coordinates
from .services.processor import UPLOADS_DIR, PageInfo, _process

# .env lives in the project root (CLAUDE.md section 8), one level above backend/.
load_dotenv(Path(__file__).resolve().parents[2] / ".env")

# A file id is uuid4().hex (see processor._process). It becomes part of a file path, so it must
# be exactly that and nothing else.
FILE_ID_RE = re.compile(r"[0-9a-f]{32}")
MAX_ERROR_DETAIL = 300

UPLOADS_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="Commitment Forms Extraction")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],
    allow_methods=["*"],
    allow_headers=["*"],
)
# Everything this app serves lives under /api. On Vercel the frontend and this backend share one
# domain: /api/* is routed here and the service receives the path unchanged (/api/upload, not
# /upload), so the routes carry the prefix themselves. Locally the frontend calls
# http://localhost:8000/api (see frontend/lib/api.ts).
API_PREFIX = "/api"
router = APIRouter(prefix=API_PREFIX)
# Layout: uploads/<file_id>/original.<ext>, uploads/<file_id>/pages/page_<n>.png, thumb_<n>.jpg
app.mount(f"{API_PREFIX}/files", StaticFiles(directory=UPLOADS_DIR), name="files")


def _public_origin(request: Request) -> str:
    """The scheme and host the browser used, e.g. https://forms.example.com.

    Behind a proxy (Vercel included) the app sees its own plain-http, internal address; the original
    one arrives in the X-Forwarded-* headers. The page-image links handed to the browser (and put in
    the Excel file) have to use it, or an https page would be given http image URLs and block them.
    """
    proto = request.headers.get("x-forwarded-proto", request.url.scheme).split(",")[0].strip()
    if proto not in ("http", "https"):
        proto = request.url.scheme
    host = (request.headers.get("x-forwarded-host") or request.headers.get("host") or request.url.netloc).split(",")[0].strip()
    return f"{proto}://{host}"


class UploadError(BaseModel):
    file_name: str
    detail: str


class UploadResponse(BaseModel):
    pages: list[PageInfo]
    errors: list[UploadError]


@router.post("/upload", response_model=UploadResponse)
async def upload(request: Request, request_files: list[UploadFile] = File(..., alias="files")):
    # _process appends /files/<id>/pages/... to this, which is where the static mount above serves them.
    base_url = f"{_public_origin(request)}{API_PREFIX}"
    if not request_files:
        raise HTTPException(status_code=400, detail="no files")
    pages: list[PageInfo] = []
    errors: list[UploadError] = []
    for f in request_files:
        try:
            pages.extend(await run_in_threadpool(_process, f, await f.read(), base_url))
        except ValueError as exc:
            errors.append(UploadError(file_name=f.filename or "file", detail=str(exc)))
    return UploadResponse(pages=pages, errors=errors)


# def _extract_fields(image_bytes: bytes) -> dict:
#     words_data = get_word_coordinates(image_bytes)
#     cropped_images = extract_field_images(image_bytes, words_data)
#     return extract_fields_with_gemini(cropped_images)


@router.post("/extract/{file_id}/{page_number}")
async def extract(file_id: str, page_number: int):
    if not FILE_ID_RE.fullmatch(file_id):
        raise HTTPException(status_code=400, detail="invalid file id")
    image_path = UPLOADS_DIR / file_id / "pages" / f"page_{page_number}.png"
    if not image_path.is_file():
        raise HTTPException(status_code=404, detail="page not found")
    image_bytes = image_path.read_bytes()
    # # Vision, OpenCV and Gemini calls all block, so the whole sequence runs off the event loop.
    # final_results = await run_in_threadpool(_extract_fields, image_bytes)
    # return final_results
    # The full page goes straight to the model - no OCR or cropping. The call blocks, so run it off
    # the event loop. {"text": ...} is the shape the frontend reads (lib/api.ts).
    # text = await run_in_threadpool(extract_full_image_with_claude, image_bytes)
    # return {"text": text}
    # The model may zoom in with crop_image over several turns before it answers.
    try:
        result = await run_in_threadpool(extract_with_agent_loop, image_bytes)
    except RuntimeError as exc:
        # The model declined, was cut off, or Gemini's quota is used up. These messages are ours
        # (see ai_service) and say what to change - nothing from the page - so they go to the table.
        raise HTTPException(status_code=502, detail=str(exc)[:MAX_ERROR_DETAIL]) from exc
    except ValueError as exc:
        raise HTTPException(status_code=502, detail="The model's reply was not valid JSON - try again") from exc
    # The frontend shows "text" as a string, so the fields go back as readable JSON text.
    return {"text": json.dumps(result, ensure_ascii=False, indent=2)}


EXPORT_COLUMNS = [
    ("first_name", "שם פרטי"),
    ("last_name", "שם משפחה"),
    ("id_number", "תעודת זהות"),
    ("central_numbers", "מספר מרכזי"),
    ("received_date", "תאריך קבלה"),
    ("sending_entity", "גורם שולח"),
    ("commitment_numbers", "מספר התחייבות"),
]
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExportRow(BaseModel):
    values: dict[str, str]
    link: str


class ExportRequest(BaseModel):
    rows: list[ExportRow]


def _build_workbook(rows: list[ExportRow]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.sheet_view.rightToLeft = True
    ws.append([label for _, label in EXPORT_COLUMNS] + ["קישור למסמך"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.values.get(key, "") for key, _ in EXPORT_COLUMNS] + [""])
        r = ws.max_row
        # Every value is text (an ID number must keep its leading zeros), and a string that
        # starts with "=" must not turn into a formula.
        for c in ws[r]:
            c.data_type = "s"
        link = ws.cell(row=r, column=len(EXPORT_COLUMNS) + 1)
        link.value = "פתיחת המסמך"
        link.hyperlink = row.link
        link.font = Font(color="0563C1", underline="single")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.post("/export")
async def export(body: ExportRequest):
    content = await run_in_threadpool(_build_workbook, body.rows)
    return Response(
        content=content,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": 'attachment; filename="commitments.xlsx"'},
    )


@router.get("/health")
def health():
    return {"status": "ok"}


# After every route above is declared: include_router copies the routes in at that moment.
app.include_router(router)
